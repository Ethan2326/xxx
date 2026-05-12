"""CSV and Excel export for LinkedIn leads."""
import io
import csv
from typing import Optional
from app.models.linkedin_lead import LinkedInLead, LinkedInCampaign
from app.services.lead_scorer import classify_lead


def leads_to_csv(leads: list[LinkedInLead]) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Nom", "Titre / Headline", "Entreprise", "Localisation",
        "LinkedIn URL", "Score", "Catégorie", "Statut",
        "Nb engagements", "Dernière activité", "Correspondance mots-clés", "Notes",
    ])
    for lead in leads:
        writer.writerow([
            lead.full_name or "",
            lead.headline or "",
            lead.company or "",
            lead.location or "",
            lead.linkedin_url,
            lead.score,
            classify_lead(lead.score),
            lead.status.value,
            lead.engagement_count,
            lead.last_engagement_at.strftime("%Y-%m-%d") if lead.last_engagement_at else "",
            "Oui" if lead.is_keyword_match else "Non",
            lead.notes or "",
        ])
    return output.getvalue().encode("utf-8-sig")  # BOM pour Excel


def leads_to_excel(leads: list[LinkedInLead], campaign_name: Optional[str] = None) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("openpyxl manquant — ajouter openpyxl aux requirements")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads LinkedIn"

    # Header
    headers = [
        "Nom", "Titre / Headline", "Entreprise", "Localisation",
        "LinkedIn URL", "Score", "Catégorie", "Statut",
        "Nb engagements", "Dernière activité", "Mots-clés", "Notes",
    ]
    header_fill = PatternFill("solid", fgColor="0A66C2")  # LinkedIn blue
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Category colors
    category_fills = {
        "hot": PatternFill("solid", fgColor="FF4444"),
        "warm": PatternFill("solid", fgColor="FF9933"),
        "cold": PatternFill("solid", fgColor="AAAAAA"),
    }

    for row_idx, lead in enumerate(leads, 2):
        category = classify_lead(lead.score)
        values = [
            lead.full_name or "",
            lead.headline or "",
            lead.company or "",
            lead.location or "",
            lead.linkedin_url,
            lead.score,
            category,
            lead.status.value,
            lead.engagement_count,
            lead.last_engagement_at.strftime("%Y-%m-%d") if lead.last_engagement_at else "",
            "Oui" if lead.is_keyword_match else "Non",
            lead.notes or "",
        ]
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if col_idx == 7:  # Catégorie column
                cell.fill = category_fills.get(category, category_fills["cold"])
                cell.font = Font(color="FFFFFF", bold=True)

    # Auto column widths
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
